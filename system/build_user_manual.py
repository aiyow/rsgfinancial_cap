"""Build manual diagrams and an editable DOCX from the canonical HTML source."""

import argparse
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from lxml import html
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
COLORS = {
    'admin': '#19634C', 'billing': '#26599B', 'resident': '#9B6515',
    'system': '#53616A', 'ink': '#15252B', 'line': '#87979D',
}
TINTS = {'admin': '#EDF6F1', 'billing': '#EEF3FA', 'resident': '#FCF5E8', 'system': '#F2F4F5'}


def font(size, bold=False):
    return ImageFont.truetype(str(Path('C:/Windows/Fonts') / ('arialbd.ttf' if bold else 'arial.ttf')), size)


def wrapped(draw, text, face, width):
    lines = []
    for paragraph in text.split('\n'):
        line = ''
        for word in paragraph.split():
            candidate = f'{line} {word}'.strip()
            if draw.textlength(candidate, font=face) > width and line:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines


def centered(draw, box, text, size=38, bold=False, color=None):
    face = font(size, bold)
    lines = wrapped(draw, text, face, box[2] - box[0] - 36)
    line_height = int(size * 1.22)
    total_height = line_height * len(lines)
    if total_height > box[3] - box[1] - 2:
        raise ValueError(f'Diagram text does not fit: {text}')
    y = (box[1] + box[3] - total_height) / 2
    for line in lines:
        draw.text(((box[0] + box[2] - draw.textlength(line, font=face)) / 2, y), line,
                  font=face, fill=color or COLORS['ink'])
        y += line_height


def arrow(draw, points, color=None, label=None, label_at=None):
    color = color or COLORS['line']
    draw.line(points, fill=color, width=5, joint='curve')
    x, y = points[-1]
    previous_x, previous_y = points[-2]
    angle = math.atan2(y - previous_y, x - previous_x)
    draw.polygon([(x, y), (x - 19 * math.cos(angle - .48), y - 19 * math.sin(angle - .48)),
                  (x - 19 * math.cos(angle + .48), y - 19 * math.sin(angle + .48))], fill=color)
    if label:
        x, y = label_at
        draw.text((x, y), label, font=font(32, True), fill=COLORS['ink'], stroke_width=6, stroke_fill='white')


def node(draw, box, title, detail, role='system', number=None):
    color = COLORS[role]
    draw.rounded_rectangle(box, radius=12, fill='white', outline=color, width=4)
    x1, y1, x2, y2 = box
    draw.rectangle((x1 + 2, y1 + 12, x1 + 10, y2 - 12), fill=color)
    if number:
        draw.text((x1 + 24, y1 + 15), number, font=font(30, True), fill=color)
        title_top = y1 + 38
    else:
        title_top = y1 + 15
    title_height = 54
    centered(draw, (x1 + 14, title_top, x2 - 10, title_top + title_height), title, 40 if number else 38, True)
    centered(draw, (x1 + 14, title_top + title_height + 2, x2 - 10, y2 - 10), detail, 36 if number else 34)


def decision(draw, box, text):
    x1, y1, x2, y2 = box
    mid_y = (y1 + y2) / 2
    draw.polygon([(x1 + 65, y1), (x2 - 65, y1), (x2, mid_y), (x2 - 65, y2),
                  (x1 + 65, y2), (x1, mid_y)], fill='#F4F6F7', outline=COLORS['line'], width=4)
    centered(draw, (x1 + 55, y1, x2 - 55, y2), text, 40, True)


def save(image, directory, name):
    target = directory / name
    image.save(target, dpi=(300, 300), optimize=True)
    print(f'Built {target.name}: {image.width} x {image.height}')


def build_assets(directory):
    directory.mkdir(parents=True, exist_ok=True)
    image = Image.new('RGB', (1800, 1660), 'white')
    draw = ImageDraw.Draw(image)
    lanes = [('ADMIN', 'admin'), ('BILLING ASSOCIATE', 'billing'), ('RESIDENT', 'resident')]
    for index, (label, role) in enumerate(lanes):
        x = 25 + index * 595
        draw.rectangle((x, 0, x + 560, 1620), fill=TINTS[role])
        centered(draw, (x, 15, x + 560, 88), label, 40, True, COLORS[role])
    steps = [
        (0, 'Prepare records', 'Units, accounts and\nresident assignments', 'admin'),
        (1, 'Import readings', 'Preview, correct and\nconfirm the workbook', 'billing'),
        (1, 'Generate and forward', 'Review the SOAs,\nthen send to Admin', 'billing'),
        (0, 'Publish bills', 'Release selected SOAs;\nemail is optional', 'admin'),
        (2, 'Pay and submit proof', 'Choose the right bill\nand payment purpose', 'resident'),
        (0, 'Verify payment', 'Approve or reject;\nupdate the balance', 'admin'),
    ]
    boxes = []
    for index, (lane, title, detail, role) in enumerate(steps):
        x, y = 45 + lane * 595, 115 + index * 250
        boxes.append((x, y, x + 520, y + 215))
    for first, second in zip(boxes, boxes[1:]):
        ax, ay = (first[0] + first[2]) / 2, first[3]
        bx, by = (second[0] + second[2]) / 2, second[1]
        middle = (ay + by) / 2
        arrow(draw, [(ax, ay), (ax, middle), (bx, middle), (bx, by)])
    for index, (box, step) in enumerate(zip(boxes, steps)):
        node(draw, box, step[1], step[2], step[3], f'{index + 1:02}')
    save(image, directory, 'monthly-cycle.png')

    image = Image.new('RGB', (1800, 1640), 'white')
    draw = ImageDraw.Draw(image)
    positions = [40, 270, 760, 990, 1220, 1450]
    data = [('Create draft', 'Coverage, due date and rates', 'billing'),
            ('Preview workbook', 'Inspect readings and warnings', 'billing'),
            ('Confirm import', 'Save the reviewed meter readings', 'billing'),
            ('Generate and review', 'Create SOAs and resolve warnings', 'billing'),
            ('Forward to Admin', 'No flagged readings left unresolved', 'billing'),
            ('Admin publishes', 'Only published SOAs reach Residents', 'admin')]
    for y1, y2 in [(200, 270), (430, 500), (650, 760), (920, 990), (1150, 1220), (1380, 1450)]:
        arrow(draw, [(540, y1), (540, y2)])
    decision(draw, (85, 500, 995, 650), 'Readings reviewed and corrected?')
    arrow(draw, [(995, 575), (1180, 575)], label='NO', label_at=(1040, 523))
    node(draw, (1180, 475, 1740, 675), 'Correct the workbook', 'Check the source log\nand preview again', 'billing')
    arrow(draw, [(1460, 475), (1460, 350), (995, 350)])
    centered(draw, (1180, 740, 1740, 910), 'Preview, import and generation\nare three separate actions.', 37)
    draw.text((563, 685), 'YES', font=font(32, True), fill=COLORS['ink'])
    for y, (title, detail, role) in zip(positions, data):
        node(draw, (85, y, 995, y + 150), title, detail, role)
    save(image, directory, 'billing-workflow.png')

    image = Image.new('RGB', (1800, 1500), 'white')
    draw = ImageDraw.Draw(image)
    data = [(35, 'Pay the correct bill', 'Resident selects the allowed purpose', 'resident'),
            (265, 'Preview the receipt', 'Check image quality and extracted facts', 'resident'),
            (495, 'Submit payment proof', 'Success message creates a pending record', 'resident'),
            (725, 'Admin checks receipt', 'Verify amount, date, method and reference', 'admin')]
    for first, second in [(195, 265), (425, 495), (655, 725), (885, 955)]:
        arrow(draw, [(820, first), (820, second)])
    decision(draw, (270, 955, 1370, 1105), 'Approve this proof?')
    arrow(draw, [(580, 1105), (580, 1160), (430, 1160), (430, 1215)], label='YES', label_at=(385, 1115))
    arrow(draw, [(1060, 1105), (1060, 1160), (1365, 1160), (1365, 1215)], label='NO', label_at=(1270, 1115))
    node(draw, (50, 1215, 830, 1460), 'Approved', 'Apply verified money to the selected SOA.\nKeep unused cash as category advance.', 'admin')
    node(draw, (960, 1215, 1740, 1460), 'Rejected', 'Read the reason. Correct the proof\nor ask Admin for clarification.', 'resident')
    arrow(draw, [(1740, 1320), (1775, 1320), (1775, 345), (1370, 345)])
    for y, title, detail, role in data:
        node(draw, (270, y, 1370, y + 160), title, detail, role)
    save(image, directory, 'payment-workflow.png')

    image = Image.new('RGB', (1800, 1500), 'white')
    draw = ImageDraw.Draw(image)
    for first, second in [(210, 285), (455, 530), (690, 790), (1000, 1080), (1250, 1330)]:
        arrow(draw, [(550, first), (550, second)])
    node(draw, (50, 40, 1050, 210), 'Resident reports an error', 'Open the published SOA and explain the issue', 'resident')
    node(draw, (50, 285, 1050, 455), 'Staff reviews the report', 'Check readings, charges and payment records', 'billing')
    decision(draw, (50, 530, 1050, 690), 'Does the SOA need correction?')
    arrow(draw, [(1050, 610), (1180, 610)], label='NO', label_at=(1080, 552))
    node(draw, (1180, 515, 1740, 725), 'Explain the result', 'Record why the original\nbill is correct', 'billing')
    arrow(draw, [(1460, 725), (1460, 1410), (1050, 1410)])
    draw.text((575, 725), 'YES', font=font(32, True), fill=COLORS['ink'])
    node(draw, (50, 790, 1050, 1000), 'Open the linked SOA correction', 'Pending or approved payment activity\nrequires an open report for that SOA.', 'billing')
    node(draw, (50, 1080, 1050, 1250), 'Save and check the revised balance', 'Keep payment identity; review category advance', 'system')
    node(draw, (50, 1330, 1050, 1490), 'Billing Associate resolves the report', 'Add a resolution note; Resident is notified', 'billing')
    save(image, directory, 'correction-workflow.png')

    image = Image.new('RGB', (1800, 480), 'white')
    draw = ImageDraw.Draw(image)
    for index, (role, label, detail) in enumerate([
        ('admin', 'ADMIN', 'Prepare and publish'),
        ('billing', 'BILLING ASSOCIATE', 'Read, bill and forward'),
        ('resident', 'RESIDENT', 'Pay and track')]):
        x = 30 + index * 600
        draw.rectangle((x, 55, x + 550, 70), fill=COLORS[role])
        centered(draw, (x, 100, x + 550, 180), label, 36, True, COLORS[role])
        centered(draw, (x, 195, x + 550, 310), detail, 39)
    centered(draw, (30, 345, 1770, 440), 'From a meter reading to a verified payment', 44, True)
    save(image, directory, 'role-overview.png')


def element(name, **attributes):
    result = OxmlElement(name)
    for key, value in attributes.items():
        result.set(qn(key), str(value))
    return result


def bookmark(paragraph, name, identifier):
    paragraph._p.insert(0, element('w:bookmarkStart', **{'w:id': identifier, 'w:name': name}))
    paragraph._p.append(element('w:bookmarkEnd', **{'w:id': identifier}))


def add_inline(paragraph, node, bold=False):
    if node.text:
        paragraph.add_run(node.text).bold = True if bold else None
    for child in node:
        if child.tag == 'a' and child.get('href', '').startswith('#'):
            hyperlink = element('w:hyperlink', **{'w:anchor': child.get('href')[1:], 'w:history': '1'})
            run = element('w:r')
            props = element('w:rPr')
            props.append(element('w:color', **{'w:val': '19634C'}))
            run.append(props)
            text = element('w:t')
            text.text = child.text_content()
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        elif child.tag == 'br':
            paragraph.add_run().add_break()
        else:
            add_inline(paragraph, child, bold or child.tag in ('strong', 'b'))
        if child.tail:
            paragraph.add_run(child.tail).bold = True if bold else None


def set_cell_shading(cell, fill):
    cell._tc.get_or_add_tcPr().append(element('w:shd', **{'w:fill': fill}))


def add_table(document, source):
    rows = source.xpath('./thead/tr|./tbody/tr|./tr')
    count = len(rows[0])
    table = document.add_table(rows=0, cols=count)
    table.autofit = False
    widths = [2.6, 3.85] if count == 2 else [2.15] * count
    if rows[0][0].text_content() in ('Role', 'History status'):
        widths = [1.45, 5.0]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = element('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(element(f'w:{edge}', **{'w:val': 'single', 'w:sz': '4', 'w:color': 'D9D9D9'}))
    table._tbl.tblPr.append(borders)
    for row_index, row_source in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(element('w:cantSplit'))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(element('w:tblHeader'))
        for col_index, (cell, cell_source) in enumerate(zip(row.cells, row_source)):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = element('w:tcMar')
            for edge in ('top', 'bottom', 'left', 'right'):
                margins.append(element(f'w:{edge}', **{'w:w': '110', 'w:type': 'dxa'}))
            cell._tc.get_or_add_tcPr().append(margins)
            if row_index == 0:
                set_cell_shading(cell, 'EAF0F4')
            elif row_index % 2 == 0:
                set_cell_shading(cell, 'F7F9FA')
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, cell_source, row_index == 0)
            for run in p.runs:
                run.font.size = Pt(10.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def render_block(document, node, source_dir):
    tag = node.tag
    classes = node.get('class', '').split()
    if tag in ('h1', 'h2', 'h3'):
        p = document.add_paragraph(style=f'Heading {int(tag[1])}')
        add_inline(p, node)
    elif tag in ('p', 'div', 'figcaption'):
        p = document.add_paragraph()
        add_inline(p, node)
        if 'audience' in classes:
            p.paragraph_format.space_after = Pt(9)
            role = node.get('data-role', 'system')
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string(COLORS[role][1:])
        elif 'meta' in classes or tag == 'figcaption':
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor.from_string('526168')
        elif 'intro' in classes:
            p.paragraph_format.space_after = Pt(12)
            for run in p.runs:
                run.font.size = Pt(12)
        elif 'subtitle' in classes:
            p.paragraph_format.space_after = Pt(18)
            for run in p.runs:
                run.font.size = Pt(17)
        elif 'eyebrow' in classes:
            p.paragraph_format.space_before = Pt(30)
            p.paragraph_format.space_after = Pt(14)
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.bold = True
    elif tag in ('ol', 'ul'):
        for index, child in enumerate(node):
            if child.tag != 'li':
                continue
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(.23)
            p.paragraph_format.first_line_indent = Inches(-.23)
            p.paragraph_format.space_after = Pt(4)
            p.add_run(f'{index + 1}. ' if tag == 'ol' else '\u2022 ')
            add_inline(p, child)
    elif tag == 'table':
        add_table(document, node)
    elif tag == 'img':
        image_path = (source_dir / node.get('src')).resolve()
        width = 1.35 if 'logo' in classes else 6.45
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if 'logo' not in classes else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = 'logo' not in classes
        picture = p.add_run().add_picture(str(image_path), width=Inches(width))
        picture._inline.docPr.set('descr', node.get('alt', ''))
    elif tag in ('figure', 'nav'):
        for child in node:
            render_block(document, child, source_dir)


def build_document(source, output):
    tree = html.parse(str(source))
    chapters = tree.xpath('//main/section')
    document = Document()
    for style in document.styles:
        for border in style._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        for fonts in style._element.xpath('./w:rPr/w:rFonts'):
            for attribute in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                fonts.attrib.pop(qn(f'w:{attribute}'), None)
        for spacing in style._element.xpath('./w:rPr/w:spacing'):
            spacing.set(qn('w:val'), '0')
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(20)
    section.header_distance = section.footer_distance = Mm(8)
    section.different_first_page_header_footer = True
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [('Title', 34), ('Heading 1', 23), ('Heading 2', 13.5), ('Heading 3', 12)]:
        style = document.styles[name]
        style.font.name, style.font.size = 'Arial', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name == 'Heading 2' else 0)
        style.paragraph_format.space_after = Pt(7 if name == 'Heading 2' else 13)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run('RSG CONDO  /  USER MANUAL')
    run.font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('RSG Condo   |   ')
    for instruction in ('PAGE', 'NUMPAGES'):
        field = element('w:fldSimple', **{'w:instr': instruction})
        field.append(element('w:r'))
        footer._p.append(field)
        if instruction == 'PAGE':
            footer.add_run(' / ')
    for run in footer.runs:
        run.font.size = Pt(8)
    for index, chapter in enumerate(chapters):
        if index:
            document.add_page_break()
        for child in chapter:
            if child.tag == 'h1':
                p = document.add_paragraph(style='Title' if index == 0 else 'Heading 1')
                add_inline(p, child)
                bookmark(p, chapter.get('id'), index + 1)
            else:
                render_block(document, child, source.parent)
    document.core_properties.title = 'RSG Condo User Manual'
    document.core_properties.subject = 'Visual workflows and role based operating instructions'
    document.core_properties.author = 'RSG Condo Project'
    document.core_properties.keywords = 'RSG Condo, workflow, manual, billing, payments'
    for paragraph in document.paragraphs:
        for border in paragraph._p.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    document.save(output)
    print(f'Built DOCX with {len(chapters)} sections: {output}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=ROOT / 'RSG_Condo_User_Manual.html')
    parser.add_argument('--output-docx', type=Path)
    parser.add_argument('--assets-only', action='store_true')
    args = parser.parse_args()
    build_assets(ROOT / 'manual_assets')
    if not args.assets_only:
        if not args.output_docx:
            parser.error('--output-docx is required unless --assets-only is used')
        build_document(args.source.resolve(), args.output_docx.resolve())


if __name__ == '__main__':
    main()
